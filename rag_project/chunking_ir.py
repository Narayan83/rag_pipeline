import re
from docx import Document
import tiktoken

enc = tiktoken.get_encoding("cl100k_base")
def ntok(t): return len(enc.encode(t))

# ---------- 1. Load docx ----------
def load_docx(path):
    doc = Document(path)
    blocks = []
    for p in doc.paragraphs:
        if p.text.strip():
            blocks.append({"type": "text", "text": p.text, "style": p.style.name})
    for i, tbl in enumerate(doc.tables):
        rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
        blocks.append({"type": "table", "table_id": i, "rows": rows})
    return blocks

# ---------- 2. Split by structure (Part / Para / OBS) ----------
PARA_RE  = re.compile(r'(?i)^\s*(?:para|part|reference number|obs[-\s]?\d+)', re.M)

def split_into_parents(blocks):
    parents, cur = [], None
    for b in blocks:
        if b["type"] == "text" and PARA_RE.match(b["text"]):
            if cur: parents.append(cur)
            cur = {"heading": b["text"], "children": []}
        else:
            if cur is None:
                cur = {"heading": "PREAMBLE", "children": []}
            cur["children"].append(b)
    if cur: parents.append(cur)
    return parents

# ---------- 3. Make child chunks ----------
MAX_TOK = 400          # child size limit
OVERLAP = 50

def chunk_text(text, max_tok=MAX_TOK, overlap=OVERLAP):
    # split on paragraphs first, then sentences if still too big
    paras = [p for p in text.split("\n") if p.strip()]
    chunks, buf = [], ""
    for p in paras:
        if ntok(buf + p) <= max_tok:
            buf += "\n" + p
        else:
            if buf: chunks.append(buf.strip())
            buf = p
            while ntok(buf) > max_tok:            # sentence fallback
                sents = re.split(r'(?<=[.!?])\s+', buf)
                head, buf = sents[0], " ".join(sents[1:])
                chunks.append(head)
    if buf: chunks.append(buf.strip())
    return chunks

def build_children(parents):
    out = []
    for pid, p in enumerate(parents):
        for c in p["children"]:
            if c["type"] == "table":
                out.append({
                    "parent_id": pid,
                    "heading":   p["heading"],
                    "type":      "table",
                    "text":      "\n".join(" | ".join(r) for r in c["rows"]),
                })
            else:
                for piece in chunk_text(c["text"]):
                    out.append({
                        "parent_id": pid,
                        "heading":   p["heading"],
                        "type":      "text",
                        "text":      piece,
                    })
    return out

# ---------- run ----------
blocks  = load_docx("IR-83EC056E63CB.docx")
parents = split_into_parents(blocks)
children = build_children(parents)

print("parents:", len(parents), "children:", len(children))